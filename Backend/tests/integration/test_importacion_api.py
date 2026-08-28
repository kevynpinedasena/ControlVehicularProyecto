"""Pruebas de integración: carga masiva de torres, apartamentos, residentes y vehículos."""
from io import BytesIO

import openpyxl
import pytest
from httpx import AsyncClient

pytestmark = pytest.mark.asyncio


def _construir_excel(hojas: dict[str, list[list[str]]]) -> bytes:
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


async def _token(cliente: AsyncClient, credenciales: dict[str, str]) -> str:
    respuesta = await cliente.post("/api/v1/auth/login", json=credenciales)
    assert respuesta.status_code == 200, respuesta.text
    return respuesta.json()["access_token"]


async def test_descargar_plantilla_devuelve_excel(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}
    respuesta = await cliente.get("/api/v1/import/template", headers=headers)
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith("application/vnd.openxmlformats")
    libro = openpyxl.load_workbook(BytesIO(respuesta.content))
    assert set(libro.sheetnames) == {
        "Torres",
        "Apartamentos",
        "Residentes",
        "Bloques",
        "Parqueaderos",
        "Vehiculos",
    }


async def test_carga_masiva_completa_crea_todo(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}
    excel = _construir_excel(
        {
            "Torres": [["nombre", "codigo"], ["Torre 9", "T9"]],
            "Apartamentos": [["torre_codigo", "numero"], ["T9", "901"]],
            "Residentes": [
                ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero", "telefono_1", "telefono_2"],
                ["Ana Ruiz", "555111", "PROPIETARIO", "T9", "901", "3000000000", ""],
            ],
            "Vehiculos": [
                ["placa", "tipo_vehiculo", "residente_documento"],
                ["XYZ111", "CARRO", "555111"],
            ],
        }
    )

    respuesta = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga.xlsx", excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200, respuesta.text
    reporte = respuesta.json()
    assert reporte["torres_creadas"] == 1
    assert reporte["apartamentos_creados"] == 1
    assert reporte["residentes_creados"] == 1
    assert reporte["vehiculos_creados"] == 1
    assert reporte["errores"] == []

    residentes = await cliente.get("/api/v1/residents?busqueda=Ana Ruiz", headers=headers)
    datos = residentes.json()
    assert len(datos) == 1
    assert datos[0]["telefono_1"] == "3000000000"
    assert datos[0]["vehiculos"][0]["placa"] == "XYZ111"


async def test_carga_masiva_reporta_errores_por_fila_sin_detener_el_resto(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}
    excel = _construir_excel(
        {
            "Torres": [
                ["nombre", "codigo"],
                ["Torre 10", "T10"],
                ["", "T10B"],  # nombre faltante -> fila inválida
            ],
            "Residentes": [
                ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero"],
                ["Sin apartamento", "777", "PROPIETARIO", "T10", "999"],  # apartamento inexistente
            ],
        }
    )

    respuesta = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga.xlsx", excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200, respuesta.text
    reporte = respuesta.json()
    assert reporte["torres_creadas"] == 1
    assert reporte["residentes_creados"] == 0
    assert len(reporte["errores"]) == 2
    assert reporte["errores"][0]["hoja"] == "Torres"
    assert reporte["errores"][1]["hoja"] == "Residentes"


async def test_carga_masiva_actualiza_registros_existentes(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}

    primera_carga = _construir_excel(
        {
            "Torres": [["nombre", "codigo"], ["Torre Original", "T20"]],
            "Apartamentos": [["torre_codigo", "numero"], ["T20", "2001"]],
            "Residentes": [
                ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero", "telefono_1"],
                ["Luis Gómez", "555222", "ARRENDATARIO", "T20", "2001", "3001112222"],
            ],
            "Vehiculos": [
                ["placa", "tipo_vehiculo", "residente_documento"],
                ["AAA111", "MOTO", "555222"],
            ],
        }
    )
    respuesta_1 = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga1.xlsx", primera_carga, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta_1.status_code == 200, respuesta_1.text
    reporte_1 = respuesta_1.json()
    assert reporte_1["torres_creadas"] == 1
    assert reporte_1["apartamentos_creados"] == 1
    assert reporte_1["residentes_creados"] == 1
    assert reporte_1["vehiculos_creados"] == 1
    assert reporte_1["errores"] == []
    assert len(reporte_1["detalles"]) == 4
    assert all(d["accion"] == "creado" for d in reporte_1["detalles"])
    assert any(d["hoja"] == "Vehiculos" and "AAA111" in d["descripcion"] for d in reporte_1["detalles"])

    # Segunda carga con las mismas claves naturales (código T20, documento 555222, placa AAA111)
    # pero datos distintos: debe actualizar en vez de fallar por duplicado.
    segunda_carga = _construir_excel(
        {
            "Torres": [["nombre", "codigo"], ["Torre Renombrada", "T20"]],
            "Apartamentos": [["torre_codigo", "numero"], ["T20", "2001"]],
            "Residentes": [
                ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero", "telefono_1"],
                ["Luis Gómez Actualizado", "555222", "ARRENDATARIO", "T20", "2001", "3009998888"],
            ],
            "Vehiculos": [
                ["placa", "tipo_vehiculo", "residente_documento"],
                ["AAA111", "CARRO", "555222"],
            ],
        }
    )
    respuesta_2 = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga2.xlsx", segunda_carga, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta_2.status_code == 200, respuesta_2.text
    reporte_2 = respuesta_2.json()
    assert reporte_2["torres_creadas"] == 0
    assert reporte_2["apartamentos_creados"] == 0
    assert reporte_2["residentes_creados"] == 0
    assert reporte_2["vehiculos_creados"] == 0
    assert reporte_2["torres_actualizadas"] == 1
    assert reporte_2["apartamentos_actualizados"] == 1
    assert reporte_2["residentes_actualizados"] == 1
    assert reporte_2["vehiculos_actualizados"] == 1
    assert reporte_2["errores"] == []
    assert len(reporte_2["detalles"]) == 4
    assert all(d["accion"] == "actualizado" for d in reporte_2["detalles"])
    assert any(
        d["hoja"] == "Residentes" and "555222" in d["descripcion"] for d in reporte_2["detalles"]
    )

    residentes = await cliente.get("/api/v1/residents?busqueda=Luis", headers=headers)
    datos = residentes.json()
    assert len(datos) == 1
    assert datos[0]["nombre"] == "Luis Gómez Actualizado"
    assert datos[0]["telefono_1"] == "3009998888"
    assert datos[0]["vehiculos"][0]["placa"] == "AAA111"
    assert datos[0]["vehiculos"][0]["tipo_vehiculo"]["codigo"] == "CARRO"


async def test_carga_masiva_crea_bloques_parqueaderos_y_asigna_vehiculo(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}
    excel = _construir_excel(
        {
            "Torres": [["nombre", "codigo"], ["Torre Parking", "TP1"]],
            "Apartamentos": [["torre_codigo", "numero"], ["TP1", "301"]],
            "Residentes": [
                ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero"],
                ["Dueño Parqueadero", "888111", "PROPIETARIO", "TP1", "301"],
            ],
            "Bloques": [["nombre", "codigo"], ["Sótano", "SOT1"]],
            "Parqueaderos": [
                ["bloque_codigo", "numero", "tipo_espacio", "tipo_vehiculo", "capacidad", "propietario_documento"],
                ["SOT1", "10", "PRIVADO", "CARRO", "1", "888111"],
                ["SOT1", "11", "COMUNAL", "MOTO", "3", ""],
            ],
            "Vehiculos": [
                ["placa", "tipo_vehiculo", "residente_documento", "bloque_codigo", "parqueadero_numero"],
                ["PRK001", "CARRO", "888111", "SOT1", "10"],
            ],
        }
    )

    respuesta = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga.xlsx", excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200, respuesta.text
    reporte = respuesta.json()
    assert reporte["bloques_creados"] == 1
    assert reporte["parqueaderos_creados"] == 2
    assert reporte["vehiculos_creados"] == 1
    assert reporte["errores"] == []
    assert any(
        d["hoja"] == "Vehiculos" and "Parqueadero SOT1/10" in d["descripcion"]
        for d in reporte["detalles"]
    )

    bloques = await cliente.get("/api/v1/parking-blocks", headers=headers)
    bloque_id = next(b["id"] for b in bloques.json() if b["codigo"] == "SOT1")

    parqueaderos = await cliente.get(f"/api/v1/parking-spots?bloque_id={bloque_id}", headers=headers)
    privado = next(p for p in parqueaderos.json() if p["numero"] == "10")
    assert privado["tipo_espacio"] == "PRIVADO"
    assert privado["propietario_residente_id"] is not None
    assert privado["vehiculos_asignados"] == 1

    vehiculo = await cliente.get("/api/v1/vehicles?busqueda=PRK001", headers=headers)
    assert vehiculo.json()[0]["parqueadero_id"] == privado["id"]


async def test_carga_masiva_parqueadero_privado_sin_propietario_reporta_error(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}
    excel = _construir_excel(
        {
            "Bloques": [["nombre", "codigo"], ["Sótano 2", "SOT2"]],
            "Parqueaderos": [
                ["bloque_codigo", "numero", "tipo_espacio", "tipo_vehiculo", "capacidad", "propietario_documento"],
                ["SOT2", "1", "PRIVADO", "CARRO", "1", ""],
            ],
        }
    )

    respuesta = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga.xlsx", excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200, respuesta.text
    reporte = respuesta.json()
    assert reporte["bloques_creados"] == 1
    assert reporte["parqueaderos_creados"] == 0
    assert len(reporte["errores"]) == 1
    assert reporte["errores"][0]["hoja"] == "Parqueaderos"


async def test_carga_masiva_vehiculo_con_parqueadero_inexistente_no_bloquea_creacion(
    cliente: AsyncClient, credenciales_super_admin: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_super_admin)}"}
    excel = _construir_excel(
        {
            "Torres": [["nombre", "codigo"], ["Torre Sin Parking", "TSP1"]],
            "Apartamentos": [["torre_codigo", "numero"], ["TSP1", "401"]],
            "Residentes": [
                ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero"],
                ["Residente Sin Parking", "888222", "PROPIETARIO", "TSP1", "401"],
            ],
            "Vehiculos": [
                ["placa", "tipo_vehiculo", "residente_documento", "bloque_codigo", "parqueadero_numero"],
                ["NOX001", "CARRO", "888222", "NOEXISTE", "1"],
            ],
        }
    )

    respuesta = await cliente.post(
        "/api/v1/import/upload",
        headers=headers,
        files={"archivo": ("carga.xlsx", excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200, respuesta.text
    reporte = respuesta.json()
    assert reporte["vehiculos_creados"] == 1
    assert len(reporte["errores"]) == 1
    assert "no existe" in reporte["errores"][0]["mensaje"]

    vehiculo = await cliente.get("/api/v1/vehicles?busqueda=NOX001", headers=headers)
    assert vehiculo.json()[0]["parqueadero_id"] is None


async def test_carga_masiva_requiere_super_admin(
    cliente: AsyncClient, credenciales_vigilante: dict[str, str]
) -> None:
    headers = {"Authorization": f"Bearer {await _token(cliente, credenciales_vigilante)}"}
    respuesta = await cliente.get("/api/v1/import/template", headers=headers)
    assert respuesta.status_code == 403
