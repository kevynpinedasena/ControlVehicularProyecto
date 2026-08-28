"""Generación de la plantilla de ejemplo (.xlsx) para la carga masiva."""
from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_FILAS_CON_VALIDACION = 200


def _agregar_hoja(
    libro: openpyxl.Workbook,
    nombre: str,
    encabezados: list[str],
    ejemplo: list[str],
    listas_desplegables: dict[str, list[str]] | None = None,
) -> None:
    hoja = libro.create_sheet(nombre)
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    hoja.append(ejemplo)
    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max(len(str(c.value)) for c in columna if c.value is not None) + 4
        hoja.column_dimensions[letra].width = ancho

    for nombre_columna, opciones in (listas_desplegables or {}).items():
        indice_columna = encabezados.index(nombre_columna) + 1
        letra_columna = get_column_letter(indice_columna)
        validacion = DataValidation(
            type="list", formula1=f'"{",".join(opciones)}"', allow_blank=False, showErrorMessage=True
        )
        validacion.error = f"Selecciona una opción de la lista: {', '.join(opciones)}"
        hoja.add_data_validation(validacion)
        validacion.add(f"{letra_columna}2:{letra_columna}{_FILAS_CON_VALIDACION}")


def generar_plantilla_excel() -> bytes:
    """Crea un libro con una hoja por entidad, encabezados y una fila de ejemplo (datos
    quemados) para que el Super Administrador sepa exactamente qué formato usar."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    _agregar_hoja(libro, "Torres", ["nombre", "codigo"], ["Torre 1", "T1"])

    _agregar_hoja(libro, "Apartamentos", ["torre_codigo", "numero"], ["T1", "101"])

    _agregar_hoja(
        libro,
        "Residentes",
        ["nombre", "documento", "tipo_residente", "torre_codigo", "apartamento_numero", "telefono_1", "telefono_2"],
        ["Juan Pérez", "123456789", "PROPIETARIO", "T1", "101", "3001234567", "3007654321"],
        listas_desplegables={"tipo_residente": ["PROPIETARIO", "ARRENDATARIO"]},
    )

    _agregar_hoja(libro, "Bloques", ["nombre", "codigo"], ["Subterráneo", "SUB"])

    _agregar_hoja(
        libro,
        "Parqueaderos",
        ["bloque_codigo", "numero", "tipo_espacio", "tipo_vehiculo", "capacidad", "propietario_documento"],
        ["SUB", "1", "PRIVADO", "CARRO", "1", "123456789"],
        listas_desplegables={
            "tipo_espacio": ["COMUNAL", "PRIVADO"],
            "tipo_vehiculo": ["CARRO", "MOTO"],
        },
    )

    _agregar_hoja(
        libro,
        "Vehiculos",
        ["placa", "tipo_vehiculo", "residente_documento", "bloque_codigo", "parqueadero_numero"],
        ["ABC123", "CARRO", "123456789", "SUB", "1"],
        listas_desplegables={"tipo_vehiculo": ["CARRO", "MOTO"]},
    )

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()